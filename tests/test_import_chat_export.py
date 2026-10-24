"""Importing WhatsApp chat exports (reports_engine/import_chat_export.py)."""

import io
import sqlite3
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path

from PIL import Image

from reports_engine import messages_db
from reports_engine.date_utils import DUBAI_TZ
from reports_engine.import_chat_export import import_exports, parse_chat, _group_for, _sender

ANDROID = """12/10/2026, 6:11 am - Messages and calls are end-to-end encrypted.
18/07/2022, 3:10 pm - Sv Mohan created group "Meydan FM Team💪"
07/10/2026, 2:12 am - Chris. Enbd: <Media omitted>
Two external lights installed above the elevators are not working and require maintenance.
08/10/2026, 6:12 pm - +971 55 123 4567: IMG-20261008-WA0001.jpg (file attached)
Q building lobby light blinking
08/10/2026, 6:20 pm - Akash: Done <This message was edited>
08/10/2026, 6:21 pm - Akash: This message was deleted
12/10/2026, 9:00 am - Akash: N building lift stuck
"""

IPHONE = """[07/10/2026, 14:05:11] Akash: Dear Team,
Kindly allow this vehicle in VIP parking slot only for today.
Plate no : H 5226
[07/10/2026, 14:06:00] Sathish: <attached: 00000012-PHOTO-2026-10-07-14-06-00.jpg>
"""


def _jpeg() -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (640, 480), (200, 30, 30)).save(out, "JPEG")
    return out.getvalue()


class TestParse(unittest.TestCase):
    def test_android_export(self):
        msgs = parse_chat(ANDROID)
        self.assertEqual([m.sender for m in msgs], ["Chris. Enbd", "+971 55 123 4567", "Akash", "Akash"])
        self.assertEqual(msgs[0].when, datetime(2026, 10, 7, 2, 12, tzinfo=DUBAI_TZ))
        self.assertTrue(msgs[0].text.startswith("Two external lights") and msgs[0].omitted_media)
        self.assertEqual((msgs[1].media, msgs[1].text), (["IMG-20261008-WA0001.jpg"], "Q building lobby light blinking"))
        self.assertEqual(msgs[2].text, "Done")                       # edit marker removed, deleted message dropped
        self.assertEqual(msgs[3].when.hour, 9)

    def test_iphone_export(self):
        msgs = parse_chat(IPHONE)
        self.assertEqual(msgs[0].when, datetime(2026, 10, 7, 14, 5, tzinfo=DUBAI_TZ))
        self.assertIn("Plate no : H 5226", msgs[0].text)
        self.assertEqual(msgs[1].media, ["00000012-PHOTO-2026-10-07-14-06-00.jpg"])

    def test_month_first_dates(self):
        msgs = parse_chat("10/13/2026, 9:00 pm - A: one\n10/14/2026, 9:00 pm - A: two\n")
        self.assertEqual([m.when.day for m in msgs], [13, 14])

    def test_group_from_file_name_and_phone_numbers(self):
        self.assertEqual(_group_for("WhatsApp Chat with Meydan FM Team_", None), "Meydan FM Team💪")
        self.assertEqual(_group_for("WhatsApp Chat with Cleaning Matters at Meydan", None), "Cleaning Matters at Meydan")
        self.assertIsNone(_group_for("WhatsApp Chat with Family", None))
        self.assertEqual(_sender("+971 55 123 4567"), "Unknown (971551234567)")
        self.assertEqual(_sender("Chris. Enbd"), "Chris. Enbd")


class TestImport(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.db = self.tmp / "whatsapp_messages.db"
        with sqlite3.connect(self.db) as c:
            messages_db.ensure_schema(c)
            # The bot started recording this group on 12-10 at 08:00
            c.execute(f"INSERT INTO messages ({messages_db.COLUMNS}) VALUES (?,?,?,?,?,?,?,?,?,?)",
                      ("BOT1", "10/12/2026, 08:00:00 AM", "Meydan FM Team💪", "Akash (1)", "first recorded", "💬 Text",
                       "", None, None, int(datetime(2026, 10, 12, 8, 0, tzinfo=DUBAI_TZ).timestamp() * 1000)))
        self.zip = self.tmp / "WhatsApp Chat with Meydan FM Team💪.zip"
        with zipfile.ZipFile(self.zip, "w") as z:
            z.writestr("WhatsApp Chat with Meydan FM Team💪.txt", ANDROID)
            z.writestr("IMG-20261008-WA0001.jpg", _jpeg())

    def tearDown(self):
        self._tmp.cleanup()

    def rows(self):
        with sqlite3.connect(self.db) as c:
            return c.execute("SELECT message_id, group_name, sender, text, media_type, thumb_jpeg IS NOT NULL, quoted_message_id "
                             "FROM messages WHERE message_id LIKE 'EXP%' ORDER BY ts_utc_ms").fetchall()

    def test_dry_run_writes_nothing(self):
        s = import_exports([self.zip], self.db)["Meydan FM Team💪"]
        self.assertEqual((s["in export"], s["already recorded by the bot"], s["new"]), (4, 1, 3))
        self.assertEqual(self.rows(), [])

    def test_imports_only_what_the_bot_did_not_record(self):
        import_exports([self.zip], self.db, apply=True)
        rows = self.rows()
        self.assertEqual(len(rows), 3)                                   # the 12-10 09:00 message was recorded by the bot
        self.assertTrue(all(r[1] == "Meydan FM Team💪" and r[6] is None for r in rows))
        self.assertEqual(rows[1][2:6], ("Unknown (971551234567)", "Q building lobby light blinking", "📷 Image", 1))
        self.assertEqual(rows[0][4], "📷 Image")                          # media left out of the export: no picture
        import_exports([self.zip], self.db, apply=True)
        self.assertEqual(len(self.rows()), 3, "re-running adds nothing")

    def test_since_filter(self):
        import_exports([self.zip], self.db, since=datetime(2026, 10, 8, tzinfo=DUBAI_TZ), apply=True)
        self.assertEqual(len(self.rows()), 2)

    def test_large_video_frame_fits_the_excel_cell(self):
        big = io.BytesIO()
        Image.new("RGB", (270, 480), (90, 90, 90)).save(big, "JPEG")   # upright phone video frame
        with sqlite3.connect(self.db) as c:
            c.execute(f"INSERT INTO messages ({messages_db.COLUMNS}) VALUES (?,?,?,?,?,?,?,?,?,?)",
                      ("VID1", "10/16/2026, 12:05:09 PM", "Meydan FM Team💪", "Akash (1)", "(Video)", "🎬 Video",
                       "", big.getvalue(), None, int(datetime(2026, 10, 16, 12, 5, tzinfo=DUBAI_TZ).timestamp() * 1000)))
        out = self.tmp / "export.xlsx"
        messages_db.export_excel(out, self.db)
        import openpyxl
        import re
        import zipfile
        ws = openpyxl.load_workbook(out).active
        self.assertEqual(ws._images, [])                                  # nothing floating over the sheet
        self.assertEqual(ws.row_dimensions[3].height, 260 * 0.75 + 4)     # VID1 is row 3 (after BOT1)
        with zipfile.ZipFile(out) as z:
            sheet = z.read("xl/worksheets/sheet1.xml").decode()
            self.assertRegex(sheet, r'<c r="G3"[^>]* t="e" vm="1"><v>#VALUE!</v></c>')   # picture inside G3
            self.assertIn("xl/metadata.xml", z.namelist())
            self.assertEqual(len([n for n in z.namelist() if re.match(r"xl/media/cellpicture\d+\.jpeg", n)]), 1)

    def test_export_is_readable(self):
        import_exports([self.zip], self.db, apply=True)
        with sqlite3.connect(self.db) as c:
            now_ms = int(datetime.now(DUBAI_TZ).timestamp() * 1000)
            c.executemany(f"INSERT INTO messages ({messages_db.COLUMNS}) VALUES (?,?,?,?,?,?,?,?,?,?)", [
                ("NEW1", "x", "Meydan FM Team💪", "Sathish (98281887682638)", "P building lift stuck", "💬 Text", "", None, None, now_ms),
                ("NEW2", "x", "Meydan FM Team💪", "Akash (1)", "@98281887682638 @971551234567 done", "💬 Text", "NEW1", None, None, now_ms + 1),
                ("NEW3", "x", "Meydan FM Team💪", "Unknown (971551234567)", "(Video)", "🎬 Video", "", None, None, now_ms + 2),
            ])
        tickets = self.tmp / "tickets.db"
        with sqlite3.connect(tickets) as c:
            c.execute("CREATE TABLE tickets (message_id TEXT, report_type TEXT, status TEXT)")
            c.execute("INSERT INTO tickets VALUES ('NEW1', 'maintenance', 'Done')")
        out = self.tmp / "export.xlsx"
        self.assertEqual(messages_db.export_excel(out, self.db, days=2, tickets_path=tickets), 3)   # older days left out
        import openpyxl
        ws = openpyxl.load_workbook(out).active
        rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rows["NEW2"][3:5], ("Akash", "@Sathish @Unknown done"))   # no numbers anywhere
        self.assertEqual(rows["NEW2"][9], "Sathish: P building lift stuck")       # reply to
        self.assertEqual(rows["NEW1"][10:12], ("Maintenance", "Done"))
        self.assertEqual((rows["NEW3"][3], rows["NEW3"][6]), ("Unknown", "No picture saved"))
        self.assertIsInstance(rows["NEW1"][8], datetime)                           # sortable time
        self.assertNotIn("971", " ".join(str(v) for r in rows.values() for v in r if v))
        full = messages_db.export_excel(out, self.db, tickets_path=tickets)
        self.assertEqual(full, 7)   # bot row, 3 imported, 3 new
        ws = openpyxl.load_workbook(out).active
        notes = {r[0]: r[6] for r in ws.iter_rows(min_row=2, values_only=True)}
        self.assertTrue(any(v == "No picture in the chat export" for k, v in notes.items() if k.startswith("EXP")))

    def test_reports_read_imported_messages(self):
        import_exports([self.zip], self.db, apply=True)
        from reports_engine.whatsapp_parser import WhatsAppParser
        with sqlite3.connect(self.db) as c:
            c.execute("INSERT OR REPLACE INTO meta (key, value) VALUES ('excel_backfill_done', '1')")
        import os
        from unittest import mock
        with mock.patch.dict(os.environ, {"REPORTS_SOURCE": "sqlite"}):
            msgs = WhatsAppParser(self.tmp / "whatsapp_messages.xlsx").load_messages(load_images=True)
        imported = [m for m in msgs if m["message_id"].startswith("EXP")]
        self.assertEqual(len(imported), 3)
        self.assertEqual(sum(1 for m in imported if m["photos_data"]), 1)


if __name__ == "__main__":
    unittest.main()
