"""Comprehensive automated test suite for Meydan Reports Automation Engine."""

import tempfile

import openpyxl
import unittest
from unittest import mock
from datetime import datetime, time, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from reports_engine.date_utils import (
    DUBAI_TZ,
    normalize_datetime_string,
    parse_datetime,
    get_shift_for_timestamp,
    is_within_freshness,
)
from reports_engine import tickets_db
from live_data import requires_live_data
from reports_engine.config_loader import get_config
from reports_engine import manager as manager_mod
from reports_engine.manager import ReportsManager, _keyword_at_word_start, _is_completion, _find_activity_to_complete, classify_operations_message, _vendor_from_recent_activity, _is_exact_closure, find_open_ticket_by_shared_word, is_closure_text
from reports_engine.repair_tickets import repair_tickets
from reports_engine.email_reader import parse_parking_email_body
from reports_engine.sender_resolver import SenderResolver
from reports_engine.deduplicator import deduplicate_events, calculate_similarity
from reports_engine.whatsapp_parser import WhatsAppParser
from reports_engine.report_generators import (
    generate_parking_report_text,
    generate_maintenance_report_text,
    generate_fire_safety_report_text,
    generate_operations_report_text,
    generate_parking_violations_report_text,
    generate_vehicle_permissions_report_text,
)


class TestDateUtils(unittest.TestCase):
    def test_arabic_numerals_normalization(self):
        raw = "١١‏/١٠‏/٢٠٢٦، ١٠:٢٩:٤٥ ص"
        norm = normalize_datetime_string(raw)
        self.assertIn("11/10/2026", norm)
        self.assertIn("AM", norm)

    def test_parse_datetime_arabic(self):
        raw = "١١‏/١٠‏/٢٠٢٦، ١٠:٢٩:٤٥ ص"
        dt = parse_datetime(raw)
        self.assertIsNotNone(dt)
        self.assertEqual(dt.year, 2026)
        self.assertEqual(dt.month, 10)
        self.assertEqual(dt.day, 11)
        self.assertEqual(dt.hour, 10)
        self.assertEqual(dt.minute, 29)

    def test_parse_datetime_western_12h(self):
        raw = "10/11/2026, 02:59:22 PM"
        dt = parse_datetime(raw)
        self.assertIsNotNone(dt)
        self.assertEqual(dt.year, 2026)
        self.assertEqual(dt.month, 10)
        self.assertEqual(dt.day, 11)
        self.assertEqual(dt.hour, 14)

    def test_parse_datetime_bot_us_format_day_le_12(self):
        dt = parse_datetime("10/12/2026, 06:12:22 AM")
        self.assertEqual((dt.month, dt.day, dt.hour), (10, 12, 6))

    def test_parse_datetime_24h_is_day_first(self):
        dt = parse_datetime("10/12/2026, 18:12")
        self.assertEqual((dt.month, dt.day), (12, 10))

    def test_shift_assignment(self):
        # Day shift: 05:00 to 16:59:59
        dt_day = datetime(2026, 10, 13, 10, 30, tzinfo=DUBAI_TZ)
        shift, s_date = get_shift_for_timestamp(dt_day)
        self.assertEqual(shift, "Day")
        self.assertEqual(s_date, "13-10-2026")

        # Night shift at 20:00
        dt_night = datetime(2026, 10, 13, 20, 0, tzinfo=DUBAI_TZ)
        shift, s_date = get_shift_for_timestamp(dt_night)
        self.assertEqual(shift, "Night")
        self.assertEqual(s_date, "13-10-2026")

        # Night shift after midnight (03:00) belongs to previous evening shift
        dt_early_morn = datetime(2026, 10, 14, 3, 0, tzinfo=DUBAI_TZ)
        shift, s_date = get_shift_for_timestamp(dt_early_morn)
        self.assertEqual(shift, "Night")
        self.assertEqual(s_date, "13-10-2026")

    def test_freshness_window(self):
        slot_dt = datetime(2026, 10, 13, 12, 0, tzinfo=DUBAI_TZ)
        # Message 30 mins before -> Fresh
        msg_30m = datetime(2026, 10, 13, 11, 30, tzinfo=DUBAI_TZ)
        self.assertTrue(is_within_freshness(msg_30m, slot_dt, window_minutes=60))

        # Message 75 mins before -> Stale
        msg_75m = datetime(2026, 10, 13, 10, 45, tzinfo=DUBAI_TZ)
        self.assertFalse(is_within_freshness(msg_75m, slot_dt, window_minutes=60))

        # Message after slot time -> Not before
        msg_after = datetime(2026, 10, 13, 12, 5, tzinfo=DUBAI_TZ)
        self.assertFalse(is_within_freshness(msg_after, slot_dt, window_minutes=60))


class TestEmailParsingAndArithmetic(unittest.TestCase):
    def setUp(self):
        self.sample_valid_email = """
MEYDAN PARKING
Date: 13-10-2026
Time: 12:00

🏢 *P BUILDING*
- Total Slots: 71
- Cars Parked: 50 🚗
- Empty Slots: 21 ✅

🏢 *N BUILDING*
- Total Slots: 77
- Cars Parked: 60 🚗
- Empty Slots: 17 ✅

🏢 *M BUILDING*
- Total Slots: 165
- Cars Parked: 100 🚗
- Empty Slots: 65 ✅

🏢 *Q BUILDING*
- Total Slots: 146
- Cars Parked: 90 🚗
- Empty Slots: 56 ✅

📊 *TOTAL SUMMARY* = 459
- Total Cars Entered: 350 ⬇️
- Total Cars Parked: 300
- Total Empty Slots: 159
"""

    def test_parse_valid_email(self):
        res = parse_parking_email_body(self.sample_valid_email)
        self.assertIsNotNone(res)
        self.assertEqual(res["date"], "13-10-2026")
        self.assertEqual(res["time"], "12:00")
        self.assertEqual(res["buildings"]["P"]["cars_parked"], 50)
        self.assertEqual(res["buildings"]["N"]["cars_parked"], 60)
        self.assertEqual(res["buildings"]["M"]["cars_parked"], 100)
        self.assertEqual(res["buildings"]["Q"]["cars_parked"], 90)
        self.assertEqual(res["summary"]["total_cars_entered"], 350)
        self.assertEqual(res["summary"]["total_cars_parked"], 300)
        self.assertEqual(res["summary"]["total_empty_slots"], 159)
        self.assertTrue(res["validation_passed"])
        self.assertEqual(len(res["discrepancies"]), 0)

    def test_email_arithmetic_discrepancy(self):
        # Tamper total parked to 320 instead of 300 (sum 50+60+100+90 = 300)
        bad_email = self.sample_valid_email.replace("Total Cars Parked: 300", "Total Cars Parked: 320")
        res = parse_parking_email_body(bad_email)
        self.assertIsNotNone(res)
        self.assertFalse(res["validation_passed"])
        self.assertTrue(any("Arithmetic mismatch" in d for d in res["discrepancies"]))


class TestSenderResolver(unittest.TestCase):
    def test_mention_resolution(self):
        resolver = SenderResolver()
        resolver.register_sender("~Md Saiful Mia 🇦🇪 (102809152299018)")
        resolver.register_sender("Masrawy (214203004629128)")

        text = "Staff car reported by @102809152299018 please contact @214203004629128"
        resolved = resolver.resolve_mentions(text)
        self.assertIn("@Md Saiful Mia", resolved)
        self.assertIn("@Masrawy", resolved)
        self.assertNotIn("@102809152299018", resolved)


class TestDeduplication(unittest.TestCase):
    def test_cross_group_duplicate_merge(self):
        dt = datetime(2026, 10, 13, 14, 0, tzinfo=DUBAI_TZ)
        events = [
            {
                "message_id": "MSG_001",
                "group_name": "Meydan FM Team💪",
                "building": "M",
                "floor": "1st Floor",
                "text": "Inside Room No. 2 AC cooling is low need to check",
                "datetime": dt,
            },
            {
                "message_id": "MSG_002",
                "group_name": "Meydan - Security Team",
                "building": "M",
                "floor": "1st Floor",
                "text": "Inside Room No 2 AC cooling is low need to check",
                "datetime": dt + timedelta(minutes=10),
            },
        ]
        unique, uncertain = deduplicate_events(events)
        self.assertEqual(len(unique), 1)
        self.assertEqual(len(uncertain), 0)
        self.assertIn("MSG_001", unique[0]["source_message_ids"])
        self.assertIn("MSG_002", unique[0]["source_message_ids"])

    def test_uncertain_near_match_flagged(self):
        dt = datetime(2026, 10, 13, 14, 0, tzinfo=DUBAI_TZ)
        events = [
            {
                "message_id": "MSG_A",
                "building": "M",
                "text": "Glass door handle loose near main entrance",
                "datetime": dt,
            },
            {
                "message_id": "MSG_B",
                "building": "M",
                "text": "Glass door not closing smoothly near entrance",
                "datetime": dt + timedelta(minutes=15),
            },
        ]
        unique, uncertain = deduplicate_events(events, similarity_threshold=0.85)
        self.assertEqual(len(unique), 2)
        self.assertTrue(len(uncertain) >= 1)


class TestReportParsersAndLinking(unittest.TestCase):
    def setUp(self):
        self.parser = WhatsAppParser()

    def test_smart_parking_parsing(self):
        sample = """
@Security
Cars entered as of 14 o'clock = 887
Cars parked as of 14 o'clock = 866

Cars parked outside = 11
"""
        res = self.parser.parse_smart_parking(sample)
        self.assertIsNotNone(res)
        self.assertEqual(res["hour_written"], 14)
        self.assertEqual(res["cars_entered"], 887)
        self.assertEqual(res["occupied_slots"], 866)
        self.assertEqual(res["empty_slots"], 887 - 866)
        self.assertEqual(res["outside_cars"], 11)

    def test_ramp_service_road_parsing(self):
        sample = """
Parking Occupany Report
Date : 13-10-2026
Time 12:00 PM
Podium Parking Total Slots : 127
Customer and visitor cars : 45
P.o.d Cars : 5
Staff Cars : 30
Total cars Entered at ramp area : 120
Total cars parked at ramp Area : 80
Total Slot Service Road : 127
Total Car Entered Service Road : 50
Total Car parked at service road : 35
Partially Blocked Parking : 2
"""
        res = self.parser.parse_ramp_service_road(sample)
        self.assertIsNotNone(res)
        self.assertEqual(res["customer_visitor"], 45)
        self.assertEqual(res["pod_cars"], 5)
        self.assertEqual(res["staff_cars"], 30)
        self.assertEqual(res["total_cars_parked_ramp"], 80)
        self.assertEqual(res["service_road_slots"], 127)
        self.assertEqual(res["service_road_parked"], 35)
        self.assertEqual(res["partially_blocked_parking"], 2)

    def test_rumble_area_parsing(self):
        sample = "Date 13th October 2026 / Rumble area location / Time09:30AM / Total car's parking at Rumble area now are 25 car's / Bus 🚌 02 car"
        res = self.parser.parse_rumble_area(sample)
        self.assertIsNotNone(res)
        self.assertEqual(res["total_cars"], 25)
        self.assertEqual(res["bus"], 2)

    def test_parking_violation_parsing(self):
        event = {
            "message_id": "V_001",
            "datetime": datetime.now(DUBAI_TZ),
            "raw_datetime": "13-10-2026 11:00 AM",
            "sender": "Ahmed (Security)",
            "group_name": "Multistory Parking Building Meydan",
            "text": "Staff parked his non-electric car on EV fast charging slot 399 on level 4 Kindly notify concern department",
            "photos": ["thumb_v1.jpg"],
        }
        res = self.parser.parse_parking_violation(event)
        self.assertIsNotNone(res)
        self.assertEqual(res["level"], "Level 4")
        self.assertEqual(res["slot_no"], "399")

    def test_vehicle_permission_parsing(self):
        event = {
            "message_id": "P_001",
            "datetime": datetime.now(DUBAI_TZ),
            "raw_datetime": "13-10-2026 09:30 AM",
            "sender": "Operations Admin (971501234567)",
            "group_name": "Parking Team at Meydan",
            "text": "Kindly allow this vehicle in smart parking / VIP parking slot. Plate no : 4 / 69888 Car make :GMC only for 3 days from today 13/10/2026 until 15/10/2026",
            "photos": [],
        }
        res = self.parser.parse_vehicle_permission(event)
        self.assertIsNotNone(res)
        self.assertEqual(res["plate_no"], "4 / 69888")
        self.assertEqual(res["car_make"], "GMC")
        self.assertIn("3 days", res["validity"])


class TestReportFormatting(unittest.TestCase):
    def test_parking_layout_reproduction(self):
        report = generate_parking_report_text(
            report_date="13-10-2026",
            report_time="14:00",
            email_data=None,
            smart_data={"total_slots": 887, "cars_entered": 500, "occupied_slots": 400, "empty_slots": 487, "outside_cars": 10},
            ramp_data={"customer_visitor": 30, "pod_cars": 2, "staff_cars": 20, "total_cars_entered_ramp": 70, "total_cars_parked_ramp": 52, "service_road_entered": 40, "service_road_parked": 25, "partially_blocked_parking": 1},
            rumble_data={"total_cars": 15, "bus": 1},
        )
        self.assertIn("🅿️ *MEYDAN PARKING*", report)
        self.assertIn("🏢 *P BUILDING*", report)
        self.assertIn("📊 *TOTAL SUMMARY* = 459", report)
        self.assertIn("🅿️ *SMART PARKING SUMMARY REPORT*", report)
        self.assertIn("🅿️ *Ramp Area Location*\n\n• Total Slots Designated : 127", report)
        self.assertIn("🛣️ *Service Road*\n\n• Total Slots : 127", report)
        self.assertIn("🏁 *Rumble Area*", report)


class TestOperationsKeywordMatching(unittest.TestCase):
    def setUp(self):
        ops = get_config()["operations"]
        self.vendors, self.activities = ops["vendors"], ops["activities"]

    def _included(self, text):
        t = text.lower()
        return bool(_keyword_at_word_start(t, self.vendors) or _keyword_at_word_start(t, self.activities))

    def test_word_forms_and_new_vendors_are_included(self):
        # Real Security Team messages from 14-10-2026 that the exact-word list missed
        self.assertTrue(self._included("These two sofa are shifted from p building to N building"))
        self.assertTrue(self._included("Al zawaya company staff remove 6 damaged chairs in m building 4th floor"))
        self.assertEqual(_keyword_at_word_start("al zawaya company staff remove", self.vendors), "al zawaya")
        self.assertEqual(_keyword_at_word_start("eurosec vendor is removing the 1 camera", self.vendors), "eurosec")
        for text in ("Table installation work done", "Chairs removed", "Boxes delivered to Q building"):
            self.assertTrue(self._included(text), text)

    def test_keywords_only_match_at_word_start(self):
        # "arki" used to match inside "parking"; bare "shift" is not a keyword (handover chatter)
        self.assertIsNone(_keyword_at_word_start("m building back side parking barrier", self.vendors))
        self.assertFalse(self._included("Night shift guard reporting, all clear"))
        self.assertFalse(self._included("Download the attendance sheet"))

    def test_guard_fault_reports_are_not_operations(self):
        # Real fault reports from 11-14 Oct 2026; they belong to maintenance, not operations
        ops = get_config()["operations"]
        faults = ops["fault_markers"]
        for text in (
            "The light near the Q Building Ground Floor near exit door is not working. Kindly inform Imdaad",
            "Two external lights installed above the elevators are not working and require maintenance.",
            "Q building first floor Speed Gate is out of service",
        ):
            t = text.lower()
            vendor = _keyword_at_word_start(t, self.vendors)
            act = _keyword_at_word_start(t, self.activities)
            self.assertIsNotNone(_keyword_at_word_start(t, faults), text)
            self.assertFalse(vendor and act, text)  # fault wins unless vendor AND activity both match
        # Vendor work on a faulty asset still counts
        t = "kone vendor installing new rope on the lift that is out of service"
        self.assertTrue(_keyword_at_word_start(t, self.vendors) and _keyword_at_word_start(t, self.activities))


# (message, expected category, expected vendor or None). Add a row whenever a keyword is added.
OPERATIONS_SCENARIOS = [
    # Real messages, 11-14 Oct 2026
    ("Al zawaya company staff remove 6 damaged chairs in m building 4th floor", "activity", "al zawaya"),
    ("These two sofa are shifted from p building to N building", "activity", None),
    ("M Building 1st Floor Group Engineering Dep, Eurosec vendor is Removing the 1 Camera.", "activity", "eurosec"),
    ("Q-Building, GF, House keeping kept office chairs and they will remove this today evening", "activity", "housekeeping"),
    ("GSU room flower watering is on going by Rose garden team", "activity", "rose garden"),
    ("Q building first floor Speed Gate is out of service", "fault", None),
    ("The light near the Q Building Ground Floor near exit door is not working. Kindly inform Imdaad", "fault", "imdaad"),
    ("Two external lights installed above the elevators are not working and require maintenance.", "fault", None),
    ("P building ground floor inside service lift one light not working need to check by technician", "fault", None),
    ("Leave Details : Full name : X Staff Id :104059 Start Date : 27 September 2027", "none", None),
    # Moving, delivery, waste
    ("Workers are loading old furniture into the truck at M building basement", "activity", None),
    ("Courier delivered 20 boxes of printer paper to Q building reception", "activity", None),
    ("Emrill staff collected the garbage bags from N building ground floor", "activity", "emrill"),
    ("IT team relocated 4 desktops from P building 2nd floor to 3rd floor", "activity", None),
    ("Supplier dropped off new water dispensers at the pantry", "activity", None),
    ("Pickup of old cartons by Farnek done", "activity", "farnek"),
    # Fit-out, civil, signage, events
    ("Contractor painting the corridor walls in M building 1st floor", "activity", None),
    ("Arki workers drilling the wall for new partition in Q building", "activity", "arki"),
    ("Hot work permit issued, welding work ongoing at service area", "activity", None),
    ("Telecon team changing branding stickers on P building lobby glass", "activity", "telecon"),
    ("Event setup in N building auditorium, chairs arranging by housekeeping", "activity", "housekeeping"),
    ("Workers removing damaged ceiling tiles in Q building corridor", "activity", None),
    # Servicing, testing, visits
    ("Kone technician repairing the lift in M building which is out of service", "activity", "kone"),
    ("Fire alarm testing by Honeywell in N building, alarms may sound", "activity", "honeywell"),
    ("DEWA meter reading visit at Q building electrical room", "activity", "dewa"),
    ("Rentokil pest control fumigation in all pantries tonight", "activity", "rentokil"),
    ("Landscaping team trimming the trees near the main gate", "activity", None),
    ("Civil defence inspection in P building, escorted by security", "activity", None),
    # Faults reported by guards
    ("Water leakage from the ceiling in M building 2nd floor washroom", "fault", None),
    ("N building main door is stuck and not closing", "fault", None),
    ("Power cut in Q building 3rd floor since 10 pm", "fault", None),
    ("AC is not cooling in the security cabin, please check", "fault", None),
    ("Glass door broken at P building entrance", "fault", None),
    ("Paint peeling off the wall near the lift lobby", "fault", None),
    ("Emergency lights are blinking in the parking area", "fault", None),
    # Chatter that must stay out
    ("Good morning team, night shift handover done, all clear", "none", None),
    ("Please everyone submit your annual leave plan", "none", None),
    ("Event in the auditorium tomorrow at 10", "none", None),
    ("Download the attendance sheet from the group", "none", None),
]

COMPLETION_SCENARIOS = [
    ("Table installation work done, vendors left.", True),
    ("Painting work completed in M building corridor", True),
    ("Arki vendor has left the site", True),
    ("Delivery finished, truck left the premises", True),
    ("Area cleared after fumigation", True),
    ("Work not yet completed, vendors will come tomorrow", False),
    ("Fumigation is pending, not done", False),
    ("Barrier on the left side not opening", False),
    ("Shift handover to night team", False),
]


class TestOperationsScenarios(unittest.TestCase):
    def setUp(self):
        self.ops = get_config()["operations"]

    def test_classification_scenarios(self):
        for text, expected_cat, expected_vendor in OPERATIONS_SCENARIOS:
            with self.subTest(text=text):
                self.assertEqual(classify_operations_message(text, self.ops), (expected_cat, expected_vendor))

    def test_completion_scenarios(self):
        for text, expected in COMPLETION_SCENARIOS:
            with self.subTest(text=text):
                self.assertEqual(_is_completion(text.lower(), self.ops["completion"]), expected)



class TestReportPeriodAndOperationsFixes(unittest.TestCase):
    def test_date_without_shift_is_full_operational_day(self):
        with mock.patch.dict("os.environ", {"REPORTS_NOW": "2026-10-15T10:00:00+04:00"}):
            self.assertEqual(ReportsManager._default_period("14-10-2026", None), ("14-10-2026", "Full"))
            self.assertEqual(ReportsManager._default_period(None, None), ("15-10-2026", "Day"))
            self.assertEqual(ReportsManager._default_period("14-10-2026", "Night"), ("14-10-2026", "Night"))
            self.assertEqual(ReportsManager._default_period("week", None), ("week", None))
        # 02:00 on the 15th: "today" still means the 14th's operational day
        with mock.patch.dict("os.environ", {"REPORTS_NOW": "2026-10-15T02:00:00+04:00"}):
            self.assertEqual(ReportsManager._default_period("15-10-2026", None), ("14-10-2026", "Full"))
        start, end = ReportsManager._get_shift_range(None, "14-10-2026", "Full")
        self.assertEqual((start.hour, start.day, end.hour, end.day), (5, 14, 5, 15))

    def test_named_location_heading_and_vendor_alias(self):
        m = ReportsManager()
        self.assertEqual(m._extract_named_location("GSU room flower watering is on going"), "GSU Room")
        aliases = get_config()["operations"]["vendor_aliases"]
        hit = _keyword_at_word_start("q-building, gf, house keeping kept office chairs", get_config()["operations"]["vendors"] + list(aliases))
        self.assertEqual(aliases.get(hit, hit), "housekeeping")

    def test_completion_detection_ignores_negation(self):
        kws = get_config()["operations"]["completion"]
        self.assertTrue(_is_completion("table installation work done, vendors left.", kws))
        self.assertTrue(_is_completion("camera removal completed", kws))
        self.assertFalse(_is_completion("work not yet completed", kws))
        self.assertFalse(_is_completion("right side barrier on the left", kws))

    def test_completion_links_to_related_open_activity(self):
        install = {"message_id": "A1", "status": "Reported", "source_message_ids": ["A1"],
                   "datetime": datetime(2026, 10, 11, 15, 1, tzinfo=DUBAI_TZ),
                   "text": "Fit Movers vendors are installing a new height-adjustable table, work is ongoing."}
        other = {"message_id": "A2", "status": "Reported", "source_message_ids": ["A2"],
                 "datetime": datetime(2026, 10, 11, 14, 59, tzinfo=DUBAI_TZ),
                 "text": "GSU room flower watering is on going by Rose garden team"}
        done = {"message_id": "C1", "quoted_message_id": "", "text": "Table installation work done, vendors left.",
                "datetime": datetime(2026, 10, 11, 16, 21, tzinfo=DUBAI_TZ)}
        target, method = _find_activity_to_complete([other, install], done)
        self.assertEqual((target["message_id"], method), ("A1", "similar"))
        # Quoted reply wins; nothing shared and no quote -> no link; older than 12h -> no link
        self.assertEqual(_find_activity_to_complete([other, install], dict(done, quoted_message_id="A2"))[0]["message_id"], "A2")
        self.assertIsNone(_find_activity_to_complete([other], done)[0])
        late = dict(done, datetime=datetime(2026, 10, 12, 4, 0, tzinfo=DUBAI_TZ))
        self.assertIsNone(_find_activity_to_complete([install], late)[0])


class TestOperationsVendorInference(unittest.TestCase):
    def _ev(self, mid, h, m, vendor, source):
        return {"message_id": mid, "datetime": datetime(2026, 10, 14, h, m, tzinfo=DUBAI_TZ),
                "vendor": vendor, "vendor_source": source}

    def test_unnamed_activity_takes_vendor_of_recent_named_one(self):
        # Real case 14-10-2026: Al Zawaya chairs at 18:08, sofa shifted (no company named) at 18:18
        drawers = self._ev("Z1", 17, 1, "Al Zawaya", "named")
        chairs = self._ev("Z2", 18, 8, "Al Zawaya", "named")
        sofa_time = datetime(2026, 10, 14, 18, 18, tzinfo=DUBAI_TZ)
        self.assertEqual(_vendor_from_recent_activity([drawers, chairs], sofa_time, 60)["message_id"], "Z2")

    def test_no_inference_when_too_old_inferred_or_disabled(self):
        old = self._ev("T1", 15, 0, "Telecon", "named")
        inferred = self._ev("I1", 18, 10, "Telecon", "inferred")
        later = self._ev("L1", 19, 0, "Arki", "named")
        t = datetime(2026, 10, 14, 18, 18, tzinfo=DUBAI_TZ)
        self.assertIsNone(_vendor_from_recent_activity([old], t, 60))          # 3h earlier
        self.assertIsNone(_vendor_from_recent_activity([inferred], t, 60))     # never chain inferences
        self.assertIsNone(_vendor_from_recent_activity([later], t, 60))        # only earlier messages
        self.assertIsNone(_vendor_from_recent_activity([self._ev("Z", 18, 8, "Al Zawaya", "named")], t, 0))


class TestMaintenanceRepliesAndPhotos(unittest.TestCase):
    def test_exact_closures_and_fault_wording(self):
        exact = get_config()["closure_exact"]
        for reply in ("Working", "Working.", "Ok now", "ok now 👍", "Fixed"):
            self.assertTrue(_is_exact_closure(reply, exact), reply)
        for text in ("Technician working on it", "lights not working properly need to check", "Ok"):
            self.assertFalse(_is_exact_closure(text, exact), text)
        # "working ..." phrases must stay out of the substring closure list (see meydan_rules.yaml)
        self.assertFalse(is_closure_text("one of the lights not working properly", get_config()["closure"]))
        self.assertFalse(any("working" in k for k in get_config()["closure"]))

    def test_missing_fault_words_are_now_faults(self):
        faults = get_config()["fault"]
        for text in ("Q building first floor Speed Gate is out of service", "Lift stuck at 3rd floor",
                     "Water leak in washroom", "Glass door broken at entrance"):
            self.assertTrue(any(k in text.lower() for k in faults), text)

    def test_reply_links_to_ticket_sharing_a_word(self):
        d = lambda h, m: datetime(2026, 10, 12, h, m, tzinfo=DUBAI_TZ)
        tickets = {
            "T1": {"message_id": "T1", "status": "Pending", "datetime": d(8, 47), "text": "Pump room hanging light is not working"},
            "T2": {"message_id": "T2", "status": "Pending", "datetime": d(9, 21), "text": "Fire exit signage light is not working"},
            "T3": {"message_id": "T3", "status": "Pending", "datetime": d(9, 30), "text": "AC is not cold in control room"},
            "T4": {"message_id": "T4", "status": "Done", "datetime": d(9, 40), "text": "Lobby light flickering"},
        }
        # Most recent OPEN ticket sharing "light"; the Done one is skipped
        self.assertEqual(find_open_ticket_by_shared_word(tickets, d(10, 59), "Light replaced")["message_id"], "T2")
        # Chatter words never link; too old -> no link
        self.assertIsNone(find_open_ticket_by_shared_word(tickets, d(10, 59), "Follow up please"))
        self.assertIsNone(find_open_ticket_by_shared_word(tickets, d(23, 59), "Light replaced"))

    def test_excel_rows_get_inline_thumbnails(self):
        import io as _io
        from PIL import Image as PILImage
        from reports_engine import excel_exporter
        buf = _io.BytesIO()
        PILImage.new("RGB", (221, 163), (200, 50, 50)).save(buf, format="JPEG")
        events = [
            {"text": "Speed gate out of service", "status": "Pending", "photos_data": [{"bytes": buf.getvalue()}]},
            {"text": "Light not working", "status": "Pending", "photos_data": []},
        ]
        with tempfile.TemporaryDirectory() as tmp, mock.patch.object(excel_exporter, "OUTPUT_DIR", Path(tmp)):
            out = excel_exporter.generate_generic_report_excel(
                report_title="t", report_text="t", events=events, unlinked_or_needs_review=[],
                source_audit=[], filename_prefix="Test",
            )
            wb = openpyxl.load_workbook(out)
            ws = wb["Events & Details"]
            self.assertEqual(ws["J1"].value, "Photo")
            self.assertEqual(len(ws._images), 1)  # only the event that has a photo
            self.assertIn("Photos", wb.sheetnames)


class TestUndatedReportsRespectChosenDate(unittest.TestCase):
    """Violations / permissions / cleaning used to ignore the Telegram date and list everything."""

    def setUp(self):
        self.m = ReportsManager()
        d = lambda day, h: datetime(2026, 10, day, h, 0, tzinfo=DUBAI_TZ)
        self.items = [{"id": "13a", "datetime": d(13, 9)}, {"id": "14a", "datetime": d(14, 20)},
                      {"id": "15early", "datetime": d(15, 3)}, {"id": "nodate", "datetime": None}]

    def test_no_date_keeps_everything(self):
        kept, label = self.m._filter_to_period(self.items, None)
        self.assertEqual((len(kept), label), (4, None))

    def test_date_keeps_only_that_operational_day(self):
        kept, label = self.m._filter_to_period(self.items, "13-10-2026")
        self.assertEqual([i["id"] for i in kept], ["13a"])
        self.assertEqual(label, "Full Day (13-10-2026)")
        # 03:00 on the 15th still belongs to the 14th's operational day (05:00 -> 05:00)
        kept, _ = self.m._filter_to_period(self.items, "14-10-2026")
        self.assertEqual([i["id"] for i in kept], ["14a", "15early"])

    def test_period_line_under_title(self):
        text = "# 🚫 Parking Violations Summary – Meydan\n\nNo active parking violations."
        out = ReportsManager._with_period_line(text, "Full Day (15-10-2026)")
        self.assertEqual(out.splitlines()[:3], ["# 🚫 Parking Violations Summary – Meydan", "", "📅 Period: Full Day (15-10-2026)"])
        self.assertEqual(ReportsManager._with_period_line(text, None), text)


class TestParkingViolationsRealDay(unittest.TestCase):
    """Multistory group, 15-10-2026: two overnight-parking violations, Akash replies "Reported"."""

    GROUP = "Multistory Parking Building Meydan"

    def _msg(self, mid, h, m, sender, text, quoted=""):
        dt = datetime(2026, 10, 15, h, m, tzinfo=DUBAI_TZ)
        return {"message_id": mid, "datetime": dt, "raw_datetime": dt.strftime("%m/%d/%Y, %I:%M:%S %p"),
                "group_name": self.GROUP, "sender": sender, "text": text, "media_type": "Text",
                "photo_preview": "", "photos_data": [], "photo_ref": "N/A", "quoted_message_id": quoted}

    def _run(self, msgs, target_date="15-10-2026"):
        m = ReportsManager()
        with mock.patch.object(m.parser, "load_messages", return_value=msgs), \
                tempfile.TemporaryDirectory() as tmp, mock.patch("reports_engine.excel_exporter.OUTPUT_DIR", Path(tmp)):
            return m.generate_parking_violations(target_date=target_date)

    def test_overnight_violations_and_reported_reply_close_ticket(self):
        msgs = [
            self._msg("V1", 3, 52, "beesyal Sunar (86535303090361)", "Smart parking staff overnight car parked at lvl 4 slot 363"),
            self._msg("V2", 7, 7, "JUNIOR (126569045606449)", "@Mohan  Staff parked her car overnight on level 4."),
            self._msg("R1", 7, 31, "Akash (219378423463939)", "Reported", quoted="V2"),
        ]
        res = self._run(msgs, target_date="14-10-2026")  # 03:52 and 07:07 on the 15th -> 14th op-day ends 05:00
        self.assertIn("📍 Level 4, Slot 363", res["text"])
        res = self._run(msgs)
        text = res["text"]
        self.assertIn("Staff parked her car overnight on level 4.", text)
        self.assertIn("✅ Status: Closed – reported to concerned department — Akash", text)
        self.assertIn("• Closed (reported to department): 1", text)
        self.assertEqual(res["closed"], 1)

    def test_unquoted_reply_links_to_latest_open_violation_within_window(self):
        msgs = [
            self._msg("V2", 7, 7, "JUNIOR", "Staff parked her car overnight on level 4."),
            self._msg("R1", 7, 40, "Akash", "Car removed"),
        ]
        self.assertIn("✅ Status: Rectified — Akash", self._run(msgs)["text"])
        late = [msgs[0], self._msg("R2", 9, 30, "Akash", "Reported")]  # 2h23m later -> not linked
        self.assertIn("⚠️ Status: Reported", self._run(late)["text"])

    def test_violation_text_mentioning_reported_is_still_a_violation(self):
        msgs = [self._msg("V3", 8, 0, "JUNIOR", "Car parked overnight on lvl 3 slot 210, reported to supervisor")]
        text = self._run(msgs)["text"]
        self.assertIn("📍 Level 3, Slot 210", text)
        self.assertIn("• Total Violations: 1", text)


def _wa_msg(mid, day, h, m, sec, group, sender, text, media="💬 Text", quoted=""):
    dt = datetime(2026, 10, day, h, m, sec, tzinfo=DUBAI_TZ)
    return {"message_id": mid, "datetime": dt, "raw_datetime": dt.strftime("%m/%d/%Y, %I:%M:%S %p"),
            "group_name": group, "sender": sender, "text": text, "media_type": media,
            "photo_preview": "", "photos_data": [], "photo_ref": "N/A", "quoted_message_id": quoted}


class TestMediaGroupingKeepsSeparateTexts(unittest.TestCase):
    def test_texts_seconds_apart_are_not_merged(self):
        # Real case 15-10 06:23: Akash sent 4 permission requests within seconds
        g, a = "Parking Team at Meydan", "Akash"
        msgs = [_wa_msg(f"P{i}", 15, 6, 23, i * 5, g, a, f"Dear Team, Kindly allow this vehicle. Plate no : X {i}")
                for i in range(4)]
        self.assertEqual(len(WhatsAppParser().group_media_events(msgs)), 4)

    def test_photo_still_attaches_to_its_text(self):
        g, s = "Fire & Safety Team Meydan", "Peace"
        msgs = [
            _wa_msg("T1", 12, 19, 39, 42, g, s, "N building 3rd floor exit light not working", media="📷 Image"),
            _wa_msg("I1", 12, 19, 39, 50, g, s, "(Image)", media="📷 Image"),
            # Real case: a second captioned photo 25s later is a separate fault, not part of the first
            _wa_msg("T2", 12, 19, 40, 7, g, s, "N building 2nd floor exit light not working", media="📷 Image"),
        ]
        evs = WhatsAppParser().group_media_events(msgs)
        self.assertEqual([e["source_message_ids"] for e in evs], [["T1", "I1"], ["T2"]])


class TestVehiclePermissionReplies(unittest.TestCase):
    G = "Parking Team at Meydan"

    def _run(self, msgs):
        m = ReportsManager()
        with mock.patch.object(m.parser, "load_messages", return_value=msgs), \
                tempfile.TemporaryDirectory() as tmp, mock.patch("reports_engine.excel_exporter.OUTPUT_DIR", Path(tmp)):
            return m.generate_vehicle_permissions()

    def test_replies_link_to_latest_request_and_counts_are_ignored(self):
        req = lambda mid, h, mi, plate: _wa_msg(mid, 13, h, mi, 0, self.G, "Akash",
                                                f"Dear Team, Kindly allow this vehicle in VIP parking slot only for today. Plate no : {plate}")
        msgs = [
            req("R1", 7, 54, "EE 80743"),
            _wa_msg("A1", 13, 8, 39, 0, self.G, "Stephen", "Parked in VIP slot @Akash"),
            req("R2", 12, 51, "AA 59845"),
            _wa_msg("A2", 13, 12, 52, 0, self.G, "Stephen", "Acknowledged"),
            _wa_msg("C1", 13, 13, 30, 0, self.G, "Saiful",
                    "*Actual Car Parked Inside at 1.30 PM In* *N: Building* *P: Building* N building total parking 0 :- 02 1 :- 0"),
        ]
        text = self._run(msgs)["text"]
        self.assertRegex(text, r"EE 80743[\s\S]*?Status: Parked in VIP slot")
        self.assertRegex(text, r"AA 59845[\s\S]*?Status: Noted")      # count report must not mark it parked
        self.assertIn("• Parked / Confirmed: 1 of 2", text)

    def test_all_noted_acknowledges_every_open_request(self):
        msgs = [_wa_msg(f"R{i}", 15, 6, 23, i * 5, self.G, "Akash",
                        f"Dear Team, Kindly allow this vehicle in the podium only for today. Plate no : M {i}")
                for i in range(3)]
        msgs.append(_wa_msg("A", 15, 7, 22, 0, self.G, "Stephen", "All Noted"))
        text = self._run(msgs)["text"]
        self.assertEqual(text.count("Status: Noted"), 3)


class TestManualCorrectionsAndOwnerKpis(unittest.TestCase):
    """Telegram correction buttons write to manual_corrections; reports must honour them."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self._env = mock.patch.dict("os.environ", {"REPORTS_DB_PATH": str(Path(self._tmp.name) / "corr.db")})
        self._env.start()

    def tearDown(self):
        self._env.stop()
        self._tmp.cleanup()

    def test_corrections_override_matching(self):
        from reports_engine.corrections_db import save_correction
        items = [
            {"message_id": "A", "status": "Pending", "source_message_ids": ["A"]},
            {"message_id": "B", "status": "Done", "resolved_by": "Raja", "source_message_ids": ["B", "B-photo"]},
            {"message_id": "C", "status": "Pending", "source_message_ids": ["C"]},
        ]
        save_correction("A", "maintenance", target_status="Done", corrected_by="owner")
        save_correction("B-photo", "maintenance", target_status="Pending", corrected_by="owner")  # via a source id
        save_correction("C", "ignore", corrected_by="owner")
        out = ReportsManager._apply_manual_corrections([dict(i) for i in items])
        self.assertEqual([(i["message_id"], i["status"]) for i in out], [("A", "Done"), ("B", "Pending")])
        self.assertEqual(out[0]["resolved_by"], "owner")
        self.assertNotIn("resolved_by", out[1])  # reopened -> no longer "done by Raja"

    @requires_live_data
    def test_report_returns_numbered_items_and_applies_correction(self):
        from reports_engine.corrections_db import save_correction
        m = ReportsManager()
        res = m.generate_maintenance(target_date="13-10-2026", target_shift="Day")
        self.assertEqual([i["no"] for i in res["items"]], list(range(1, len(res["items"]) + 1)))
        in_shift = next(t for t in m.last_items["maintenance"] if not t.get("is_carry_over") and t["status"] == "Pending")
        carried = next(t for t in m.last_items["maintenance"] if t.get("is_carry_over") and t["status"] == "Pending")
        save_correction(in_shift["message_id"], "maintenance", target_status="Done", corrected_by="owner")
        save_correction(carried["message_id"], "maintenance", target_status="Done", corrected_by="owner")
        res2 = m.generate_maintenance(target_date="13-10-2026", target_shift="Day")
        ids2 = {i["id"]: i["status"] for i in res2["items"]}
        self.assertEqual(ids2[in_shift["message_id"]], "Done")
        self.assertNotIn(carried["message_id"], ids2)  # a fixed old ticket no longer carries over

    @requires_live_data
    def test_owner_kpis_sections(self):
        from reports_engine.owner_kpis import generate_owner_kpis
        text = generate_owner_kpis(ReportsManager())["text"]
        for section in ("الصيانة", "الحريق والسلامة", "العمليات", "مخالفات المواقف", "تصاريح المركبات", "SLA"):
            self.assertIn(section, text)


class TestSqliteMessageStore(unittest.TestCase):
    """The SQLite store must give the parser exactly what the Excel workbook gives it."""

    SRC = Path(__file__).resolve().parent.parent / "workspace" / "whatsapp_messages.xlsx"

    def setUp(self):
        if not self.SRC.exists():
            self.skipTest("live workspace/whatsapp_messages.xlsx not present")
        import shutil
        self._tmp = tempfile.TemporaryDirectory()
        self.xlsx = Path(self._tmp.name) / "whatsapp_messages.xlsx"
        shutil.copy(self.SRC, self.xlsx)
        self._env = mock.patch.dict("os.environ", {"REPORTS_SOURCE": "auto"})
        self._env.start()

    def tearDown(self):
        self._env.stop()
        self._tmp.cleanup()

    def test_excel_used_until_import_completes(self):
        from reports_engine import messages_db
        db = messages_db.db_path_for(self.xlsx)
        self.assertEqual(WhatsAppParser(self.xlsx)._message_source()[0], "excel")
        # A database the bot started writing before the import is not trusted yet
        import sqlite3
        with sqlite3.connect(db) as c:
            c.executescript(messages_db.SCHEMA)
        self.assertFalse(messages_db.is_ready(db))
        self.assertEqual(WhatsAppParser(self.xlsx)._message_source()[0], "excel")

    def test_import_matches_excel_exactly(self):
        from reports_engine import messages_db
        r = messages_db.import_excel(self.xlsx)
        self.assertGreater(r["total"], 0)
        self.assertGreater(r["with_photo"], 0)
        self.assertEqual(WhatsAppParser(self.xlsx)._message_source()[0], "sqlite")
        c = messages_db.compare(self.xlsx)
        self.assertEqual((c["missing_in_sqlite"], c["only_in_sqlite"], c["field_differences"]), (0, 0, 0))
        # Re-running is safe and keeps the same data
        self.assertEqual(messages_db.import_excel(self.xlsx)["total"], r["total"])

    def test_reports_identical_from_both_sources(self):
        from reports_engine import messages_db
        messages_db.import_excel(self.xlsx)
        texts = {}
        for source in ("excel", "sqlite"):
            with mock.patch.dict("os.environ", {"REPORTS_SOURCE": source}):
                m = ReportsManager(self.xlsx)
                texts[source] = (m.generate_maintenance(target_date="13-10-2026", target_shift="Day")["text"],
                                 m.generate_operations(target_date="week")["text"])
        self.assertEqual(texts["excel"], texts["sqlite"])


@requires_live_data
class TestOwnerDigestAndOverdue(unittest.TestCase):
    def test_overdue_lists_only_open_tickets_older_than_48h(self):
        from reports_engine.owner_digest import generate_overdue
        res = generate_overdue(ReportsManager(), now=datetime(2026, 10, 14, 9, 0, tzinfo=DUBAI_TZ))
        items = res["overdue"]
        self.assertEqual(res["total"], len(items))
        self.assertTrue(all(i["hours_open"] > 48 for i in items))
        self.assertTrue(all(i["report"] in ("maintenance", "fire_safety") for i in items))
        self.assertEqual([i["hours_open"] for i in items], sorted((i["hours_open"] for i in items), reverse=True))
        self.assertEqual(len({i["id"] for i in items}), len(items))

    def test_overdue_threshold_and_corrections(self):
        from reports_engine.owner_digest import generate_overdue
        from reports_engine.corrections_db import save_correction
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.dict("os.environ", {"REPORTS_DB_PATH": str(Path(tmp) / "corr.db")}):
            m = ReportsManager()
            now = datetime(2026, 10, 14, 9, 0, tzinfo=DUBAI_TZ)
            first = generate_overdue(m, now=now)["overdue"]
            self.assertGreater(len(first), 0)
            # A ticket marked fixed from the reminder's ✅ button is no longer overdue
            save_correction(first[0]["id"], first[0]["report"], target_status="Done", corrected_by="owner")
            after = generate_overdue(m, now=now)["overdue"]
            self.assertNotIn(first[0]["id"], [i["id"] for i in after])
            self.assertEqual(generate_overdue(m, min_hours=10_000, now=now)["overdue"], [])

    def test_morning_digest_sections(self):
        from reports_engine.owner_digest import generate_morning_digest
        text = generate_morning_digest(ReportsManager(), now=datetime(2026, 10, 14, 8, 0, tzinfo=DUBAI_TZ))["text"]
        for part in ("ملخص الصباح", "الصيانة", "الحريق والسلامة", "مخالفات المواقف", "تصاريح اليوم", "مفتوحة الآن"):
            self.assertIn(part, text)


class TestPrivacyAndMasterData(unittest.TestCase):
    def test_ts_hr_markers_match_config(self):
        import re as _re
        ts = (Path(__file__).resolve().parent.parent / "src" / "privacy.ts").read_text()
        ts_markers = _re.findall(r"'([^']+)'", ts.split("HR_MARKERS = [")[1].split("];")[0])
        self.assertEqual(ts_markers, get_config()["exclude"]["leave_markers"])

    def test_purge_hr_respects_retention(self):
        import sqlite3
        from reports_engine import messages_db
        with tempfile.TemporaryDirectory() as tmp:
            db = Path(tmp) / "m.db"
            old_ms = int((datetime.now(DUBAI_TZ) - timedelta(days=400)).timestamp() * 1000)
            new_ms = int(datetime.now(DUBAI_TZ).timestamp() * 1000)
            with sqlite3.connect(db) as c:
                messages_db.ensure_schema(c)
                c.executemany(
                    "INSERT INTO messages (message_id, timestamp_text, group_name, text, ts_utc_ms) VALUES (?,?,?,?,?)",
                    [("hr-old", "", "g", "Leave Details : Full name : X", old_ms),
                     ("hr-new", "", "g", "Leave Details : Full name : Y", new_ms),
                     ("fault-old", "", "g", "Q building speed gate out of service", old_ms)])
            markers = get_config()["exclude"]["leave_markers"]
            self.assertEqual(messages_db.purge_hr(db, None, markers), 0)  # default: keep forever
            self.assertEqual(messages_db.purge_hr(db, 365, markers), 1)
            with sqlite3.connect(db) as c:
                left = {r[0] for r in c.execute("SELECT message_id FROM messages")}
            self.assertEqual(left, {"hr-new", "fault-old"})

    def test_display_name_hides_phone_numbers(self):
        from reports_engine import master_data
        self.assertEqual(master_data.display_name("Shrikant Yadav (209852118868213)"), "Shrikant Yadav")
        with mock.patch.object(master_data, "load", return_value={"staff": {"209852118868213": {"name": "Eng. Shrikant", "team": "FM"}}}):
            self.assertEqual(master_data.display_name("Shrikant Yadav (209852118868213)"), "Eng. Shrikant")
        self.assertEqual(master_data.display_name(None), "")

    def test_building_aliases(self):
        from reports_engine import master_data
        self.assertEqual(master_data.building_from_aliases("Light not working at Q blk ground floor"), "Q")
        self.assertIsNone(master_data.building_from_aliases("Lobby light not working"))
        bld, _, heading = WhatsAppParser().extract_building_and_floor("Light not working at q-blk 2nd floor")
        self.assertEqual((bld, heading), ("Q", "Building Q"))


class TestSlaHiddenFromReports(unittest.TestCase):
    def test_sla_breach_not_rendered_in_report_text(self):
        item = {
            "text": "N building emergency exit light is not working",
            "datetime": datetime(2026, 10, 12, 19, 39, tzinfo=DUBAI_TZ),
            "status": "Pending",
            "hours_open": 50.0,
            "sla_breach": True,
            "location_heading": "Building N",
        }
        stats = {"completed": 0, "total": 1, "open": 1, "sla_breaches": 1,
                 "safety_priority": "N/A", "teams": "N/A", "shift_info": "Day Shift (14-10-2026)"}
        for render in (generate_maintenance_report_text, generate_fire_safety_report_text):
            text = render([dict(item)], dict(stats))
            self.assertNotIn("SLA", text, render.__name__)
            self.assertIn("🕐 Open: 2.1d", text, render.__name__)


class TestTicketsSla(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.db_path = str(Path(self._tmp.name) / "tickets.db")
        tickets_db.init_tickets_table(self.db_path)

    def tearDown(self):
        self._tmp.cleanup()

    def test_hours_open_uses_parsed_datetime_not_raw_text(self):
        created = datetime(2026, 10, 12, 6, 12, 22, tzinfo=DUBAI_TZ)
        now_iso = (created + timedelta(hours=30)).isoformat()
        ticket = {
            "message_id": "msg-sla-1",
            "datetime": created,
            "raw_datetime": "10/12/2026, 06:12:22 AM",
            "status": "Pending",
            "text": "AC not working",
        }
        tickets_db.enrich_ticket(ticket, "maintenance", "FM", "Morning", now_iso, db_path=self.db_path)

        self.assertGreater(ticket["hours_open"], 0)
        self.assertAlmostEqual(ticket["hours_open"], 30.0, places=1)
        self.assertTrue(ticket["sla_breach"])
        stored = tickets_db.get_first_seen("msg-sla-1", self.db_path)
        self.assertEqual(datetime.fromisoformat(stored), created)

    def test_legacy_non_iso_first_seen_is_reparsed_and_repaired(self):
        legacy_raw = "١١‏/١٠‏/٢٠٢٦، ١٠:٢٩:٤٥ ص"  # 11 Oct 2026 10:29:45 (DD/MM)
        tickets_db.upsert_ticket(
            "msg-sla-2", "maintenance", "FM", None, None, "Leak", "Pending", None,
            first_seen=legacy_raw, last_seen=legacy_raw, shift_label="Morning", db_path=self.db_path,
        )
        now_iso = datetime(2026, 10, 12, 16, 29, 45, tzinfo=DUBAI_TZ).isoformat()
        ticket = {"message_id": "msg-sla-2", "status": "Pending", "text": "Leak"}
        tickets_db.enrich_ticket(ticket, "maintenance", "FM", "Morning", now_iso, db_path=self.db_path)

        self.assertAlmostEqual(ticket["hours_open"], 30.0, places=1)
        self.assertTrue(ticket["sla_breach"])
        stored = tickets_db.get_first_seen("msg-sla-2", self.db_path)
        self.assertEqual(datetime.fromisoformat(stored), datetime(2026, 10, 11, 10, 29, 45, tzinfo=DUBAI_TZ))

    def test_repair_script_rewrites_legacy_rows(self):
        iso = datetime(2026, 10, 12, 9, 0, tzinfo=DUBAI_TZ).isoformat()
        rows = {
            "msg-us": "10/12/2026, 06:12:22 AM",      # en-US MM/DD -> 12 Oct
            "msg-ar": "١١‏/١٠‏/٢٠٢٦، ١٠:٢٩:٤٥ ص",     # Arabic DD/MM -> 11 Oct
            "msg-iso": iso,
            "msg-bad": "not a date",
        }
        for mid, fs in rows.items():
            tickets_db.upsert_ticket(
                mid, "maintenance", "FM", None, None, "x", "Pending", None,
                first_seen=fs, last_seen=iso, shift_label="Morning", db_path=self.db_path,
            )

        dry = repair_tickets(self.db_path, apply=False)
        self.assertEqual({e["message_id"] for e in dry["fixed"]}, {"msg-us", "msg-ar"})
        self.assertEqual([e["message_id"] for e in dry["unparseable"]], ["msg-bad"])
        self.assertEqual(tickets_db.get_first_seen("msg-us", self.db_path), rows["msg-us"])  # dry run wrote nothing

        repair_tickets(self.db_path, apply=True)
        get = lambda mid: datetime.fromisoformat(tickets_db.get_first_seen(mid, self.db_path))
        self.assertEqual(get("msg-us"), datetime(2026, 10, 12, 6, 12, 22, tzinfo=DUBAI_TZ))
        self.assertEqual(get("msg-ar"), datetime(2026, 10, 11, 10, 29, 45, tzinfo=DUBAI_TZ))
        self.assertEqual(tickets_db.get_first_seen("msg-iso", self.db_path), iso)
        self.assertEqual(tickets_db.get_first_seen("msg-bad", self.db_path), "not a date")
        self.assertEqual(repair_tickets(self.db_path)["fixed"], [])  # idempotent


if __name__ == "__main__":
    unittest.main()


@requires_live_data
class TestDataQualityReport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from reports_engine.manager import ReportsManager
        from reports_engine.data_quality import generate_data_quality
        cls.res = generate_data_quality(ReportsManager())

    def test_metrics_are_internally_consistent(self):
        m = self.res["metrics"]
        self.assertEqual(sum(m["per_group"].values()), m["messages"])
        self.assertLessEqual(m["reply_used"], m["reply_captured"])
        self.assertLessEqual(m["reply_captured"], m["messages"])
        self.assertLessEqual(m["tickets_without_building"], m["tickets"])
        self.assertLessEqual(m["permissions_without_plate"], m["permissions"])

    def test_reply_warning_needs_recorded_reply_info(self):
        m = self.res["metrics"]
        if m["messages"] and m["reply_captured"] / m["messages"] < 0.8:
            self.assertNotIn("أكتر من الربط بالـ Reply", self.res["text"])
            self.assertIn("مش دقيقة لسه", self.res["text"])
