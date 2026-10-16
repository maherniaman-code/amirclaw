"""Excel workbook generation for Meydan reports using openpyxl."""

import io
import time
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .config_loader import BASE_DIR
from .date_utils import format_report_timestamp

OUTPUT_DIR = BASE_DIR / "workspace" / "downloads" / "reports"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
REGULAR_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", size=10, bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)


def apply_table_styling(ws, header_row_idx: int = 1) -> None:
    ws.views.sheetView[0].showGridLines = True
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        is_zebra = (row_idx % 2 == 1)
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if is_zebra:
                cell.fill = ZEBRA_FILL
            cell.border = THIN_BORDER
            if cell.alignment.horizontal is None:
                cell.alignment = Alignment(vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = min(len(val_str), 60)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def generate_parking_excel(
    report_text: str,
    email_data: Optional[Dict[str, Any]],
    smart_data: Optional[Dict[str, Any]],
    ramp_data: Optional[Dict[str, Any]],
    rumble_data: Optional[Dict[str, Any]],
    source_audit: List[Dict[str, Any]],
    filename_prefix: str = "Parking_Report",
) -> Path:
    wb = openpyxl.Workbook()
    # 1. Parking Summary
    ws_summary = wb.active
    ws_summary.title = "Parking Summary"
    # Plain text without asterisks
    clean_text = report_text.replace("*", "")
    for line in clean_text.split("\n"):
        ws_summary.append([line])
    ws_summary.column_dimensions["A"].width = 65

    # 2. Building Details
    ws_bld = wb.create_sheet("Building Details")
    ws_bld.append(["Building", "Total Slots", "Cars Parked", "Empty Slots", "Status"])
    bld_data = email_data.get("buildings", {}) if email_data else {}
    for b in ["P", "N", "M", "Q"]:
        b_info = bld_data.get(b, {})
        ws_bld.append([
            f"{b} Building",
            b_info.get("total_slots", "N/A"),
            b_info.get("cars_parked", "N/A"),
            b_info.get("empty_slots", "N/A"),
            "Verified" if email_data and email_data.get("validation_passed") else "Unverified",
        ])
    apply_table_styling(ws_bld)

    # 3. Smart Parking
    ws_smart = wb.create_sheet("Smart Parking")
    ws_smart.append(["Total Slots", "Cars Entered", "Occupied Slots", "Empty Slots", "Outside Cars"])
    smart = smart_data or {}
    ws_smart.append([
        smart.get("total_slots", 887),
        smart.get("cars_entered", "N/A"),
        smart.get("occupied_slots", "N/A"),
        smart.get("empty_slots", "N/A"),
        smart.get("outside_cars", "N/A"),
    ])
    apply_table_styling(ws_smart)

    # 4. Ramp Area
    ws_ramp = wb.create_sheet("Ramp Area")
    ws_ramp.append(["Metric", "Value"])
    ramp = ramp_data or {}
    ws_ramp.append(["Total Slots Designated", ramp.get("podium_slots", 127)])
    ws_ramp.append(["Customer and Visitors", ramp.get("customer_visitor", "N/A")])
    ws_ramp.append(["P.O.D", ramp.get("pod_cars", "N/A")])
    ws_ramp.append(["Staff", ramp.get("staff_cars", "N/A")])
    ws_ramp.append(["Total car Entered", ramp.get("total_cars_entered_ramp", "N/A")])
    ws_ramp.append(["Total Car Parking at Ramp Area", ramp.get("total_cars_parked_ramp", "N/A")])
    apply_table_styling(ws_ramp)

    # 5. Service Road
    ws_sr = wb.create_sheet("Service Road")
    ws_sr.append(["Metric", "Value", "Notes"])
    ws_sr.append(["Total Slots", ramp.get("service_road_slots", 127), "As sent by the team (127 when left out)"])
    ws_sr.append(["Cars Entered", ramp.get("service_road_entered", "N/A"), ""])
    ws_sr.append(["Total Occupied Slots", ramp.get("service_road_parked", "N/A"), ""])
    ws_sr.append(["Partially Blocked Parking", ramp.get("partially_blocked_parking", "N/A"), ""])
    apply_table_styling(ws_sr)

    # 6. Rumble Area
    ws_rumble = wb.create_sheet("Rumble Area")
    ws_rumble.append(["Date Written", "Time Written", "Total Cars", "Bus"])
    rumble = rumble_data or {}
    ws_rumble.append([
        rumble.get("date_written", "N/A"),
        rumble.get("time_written", "N/A"),
        rumble.get("total_cars", "N/A"),
        rumble.get("bus", "N/A"),
    ])
    apply_table_styling(ws_rumble)

    # 7. Source Audit
    ws_audit = wb.create_sheet("Source Audit")
    ws_audit.append(["Source ID", "Group Name", "Sender", "Timestamp", "Extracted Section", "Validation Status"])
    for a in source_audit:
        ws_audit.append([
            a.get("source_id", "N/A"),
            a.get("group_name", "N/A"),
            a.get("sender", "N/A"),
            format_report_timestamp(a.get("timestamp")),
            a.get("section", "N/A"),
            a.get("status", "N/A"),
        ])
    apply_table_styling(ws_audit)

    out_file = OUTPUT_DIR / f"{filename_prefix}_{int(time.time())}.xlsx"
    wb.save(str(out_file))
    return out_file


def _make_thumbnail(raw_bytes: bytes, max_side: int, quality: int) -> "Tuple[bytes, int, int]":
    """Compress image to JPEG thumbnail. Returns (jpeg_bytes, width_px, height_px)."""
    from PIL import Image as PilImage
    img = PilImage.open(io.BytesIO(raw_bytes)).convert("RGB")
    w, h = img.size
    if w > max_side or h > max_side:
        scale = min(max_side / w, max_side / h)
        img = img.resize((max(1, int(w * scale)), max(1, int(h * scale))), PilImage.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=quality, optimize=True)
    return buf.getvalue(), img.width, img.height


def generate_violations_excel(
    report_title: str,
    report_text: str,
    events: List[Dict[str, Any]],
    unlinked_or_needs_review: List[Dict[str, Any]],
    source_audit: List[Dict[str, Any]],
    filename_prefix: str,
    thumbnail_max_side: int = 160,
    thumbnail_quality: int = 70,
) -> Dict[str, Any]:
    """
    Report 5 — Parking Violations Excel with per-row embedded JPEG thumbnails.
    Returns dict with excel_path, file_size_bytes, photo_mapping.
    Never embeds an image that belongs to another row.
    """
    from openpyxl.drawing.image import Image as OpenpyxlImage

    wb = openpyxl.Workbook()

    # ── Sheet 1: Report Summary ────────────────────────────────────
    ws_summary = wb.active
    ws_summary.title = "Report Summary"
    for line in report_text.split("\n"):
        ws_summary.append([line])
    ws_summary.column_dimensions["A"].width = 70

    # ── Sheet 2: Events & Details (with Photo column) ──────────────
    ws_events = wb.create_sheet("Events & Details")
    headers = [
        "Event #", "Level / Slot", "Description", "Status",
        "Sender", "Timestamp", "Source Message IDs", "Quoted Message ID", "Photo",
    ]
    ws_events.append(headers)

    THUMB_ROW_HEIGHT = max(90, int(thumbnail_max_side * 0.75) + 10)
    PHOTO_COL = 9  # column I

    photo_mapping: List[Dict[str, Any]] = []

    for idx, ev in enumerate(events, 1):
        data_row = idx + 1  # row 1 is header
        ws_events.append([
            idx,
            ev.get("location_heading") or ev.get("level") or "N/A",
            ev.get("issue") or ev.get("activity") or ev.get("text") or "N/A",
            ev.get("status", "Reported"),
            ev.get("sender", "N/A"),
            format_report_timestamp(ev.get("datetime") or ev.get("raw_datetime")),
            ", ".join(ev.get("source_message_ids", [])),
            ev.get("quoted_message_id", "") or "None",
            "",
        ])

        photos = ev.get("photos_data", [])
        if photos:
            p_info = photos[0]  # only the image mapped to this event's own row
            raw = p_info.get("bytes")
            media_path = p_info.get("media_path", "unknown")
            if raw:
                try:
                    thumb_bytes, tw, th = _make_thumbnail(raw, thumbnail_max_side, thumbnail_quality)
                    img = OpenpyxlImage(io.BytesIO(thumb_bytes))
                    img.width = tw
                    img.height = th
                    col_letter = get_column_letter(PHOTO_COL)
                    ws_events.add_image(img, f"{col_letter}{data_row}")
                    ws_events.row_dimensions[data_row].height = THUMB_ROW_HEIGHT
                    photo_mapping.append({
                        "excel_row": data_row,
                        "message_id": ev.get("message_id", "N/A"),
                        "image_file": Path(media_path).name,
                        "thumb_px": f"{tw}×{th}",
                    })
                except Exception:
                    ws_events.cell(row=data_row, column=PHOTO_COL).value = "N/A"
            else:
                ws_events.cell(row=data_row, column=PHOTO_COL).value = "N/A"
        else:
            ws_events.cell(row=data_row, column=PHOTO_COL).value = "N/A"

    ws_events.column_dimensions[get_column_letter(PHOTO_COL)].width = 24
    apply_table_styling(ws_events)

    # ── Sheet 3: Photos ────────────────────────────────────────────
    events_with_photos = [e for e in events if e.get("photos_data")]
    if events_with_photos:
        ws_photos = wb.create_sheet("Photos")
        ws_photos.append(["Photo #", "Event #", "Level / Slot", "Image"])
        for p_idx, ev in enumerate(events_with_photos, 1):
            r_num = p_idx + 1
            ev_num = next((m["excel_row"] - 1 for m in photo_mapping if m["message_id"] == ev.get("message_id")), "N/A")
            ws_photos.append([p_idx, ev_num, ev.get("location_heading") or ev.get("level") or "N/A", ""])
            raw = ev["photos_data"][0].get("bytes")
            if raw:
                try:
                    thumb_bytes, tw, th = _make_thumbnail(raw, thumbnail_max_side, thumbnail_quality)
                    img = OpenpyxlImage(io.BytesIO(thumb_bytes))
                    img.width = tw
                    img.height = th
                    ws_photos.add_image(img, f"D{r_num}")
                    ws_photos.row_dimensions[r_num].height = THUMB_ROW_HEIGHT
                except Exception:
                    pass
        ws_photos.column_dimensions["D"].width = 24
        apply_table_styling(ws_photos)

    # ── Sheet 4: Needs Review & Unlinked ──────────────────────────
    ws_review = wb.create_sheet("Needs Review & Unlinked")
    ws_review.append(["Item ID", "Type / Reason", "Message Text", "Sender", "Timestamp"])
    for r in unlinked_or_needs_review:
        r = _review_row(r)
        ws_review.append([
            r.get("message_id", "N/A"),
            r.get("reason", "Unlinked reply"),
            r.get("text", "N/A"),
            r.get("sender", "N/A"),
            format_report_timestamp(r.get("datetime") or r.get("raw_datetime")),
        ])
    apply_table_styling(ws_review)

    # ── Sheet 5: Source Audit ──────────────────────────────────────
    ws_audit = wb.create_sheet("Source Audit")
    ws_audit.append(["Message ID", "Group", "Sender", "Timestamp", "Quoted Message ID", "Audit Status"])
    for a in source_audit:
        ws_audit.append([
            a.get("message_id", "N/A"),
            a.get("group_name", "N/A"),
            a.get("sender", "N/A"),
            format_report_timestamp(a.get("timestamp") or a.get("raw_datetime")),
            a.get("quoted_message_id", "") or "Empty",
            a.get("audit_status", "Valid"),
        ])
    apply_table_styling(ws_audit)

    out_file = OUTPUT_DIR / f"{filename_prefix}_{int(time.time())}.xlsx"
    wb.save(str(out_file))
    file_size = out_file.stat().st_size

    return {
        "excel_path": out_file,
        "file_size_bytes": file_size,
        "photo_mapping": photo_mapping,
    }


def _review_row(r: Dict[str, Any]) -> Dict[str, Any]:
    """Flatten deduplicator flags ({type, current_event/earlier_event | event_a/event_b}) into a review row."""
    if "type" not in r:
        return r
    later = r.get("current_event") or r.get("event_b") or {}
    earlier = r.get("earlier_event") or r.get("event_a") or {}
    pct = f"{r['similarity']:.0%} similar" if isinstance(r.get("similarity"), (int, float)) else "similar"
    when = format_report_timestamp(earlier.get("datetime") or earlier.get("raw_datetime"))
    return {
        "message_id": later.get("message_id", "N/A"),
        "reason": f"{r['type']} ({pct} to {earlier.get('message_id', 'N/A')} at {when})",
        "text": later.get("text", "N/A"),
        "sender": later.get("sender", "N/A"),
        "datetime": later.get("datetime"),
        "raw_datetime": later.get("raw_datetime"),
    }


def _embed_row_thumbnails(ws, events: List[Dict[str, Any]], col_letter: str, first_row: int, max_px: int = 80) -> None:
    """Small inline preview of each event's first photo (WhatsApp thumbnails are already compressed)."""
    from openpyxl.drawing.image import Image as OpenpyxlImage
    has_any = False
    for offset, ev in enumerate(events):
        photos = ev.get("photos_data") or []
        raw = photos[0].get("bytes") if photos else None
        if not raw:
            continue
        try:
            img = OpenpyxlImage(io.BytesIO(raw))
            scale = min(1.0, max_px / max(img.width, img.height))
            img.width, img.height = int(img.width * scale), int(img.height * scale)
            row = first_row + offset
            ws.add_image(img, f"{col_letter}{row}")
            ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 15, img.height * 0.75 + 6)
            has_any = True
        except Exception:
            continue
    if has_any:
        ws.column_dimensions[col_letter].width = max_px / 7 + 2


def generate_generic_report_excel(
    report_title: str,
    report_text: str,
    events: List[Dict[str, Any]],
    unlinked_or_needs_review: List[Dict[str, Any]],
    source_audit: List[Dict[str, Any]],
    filename_prefix: str,
) -> Path:
    wb = openpyxl.Workbook()
    # 1. Report Summary
    ws_summary = wb.active
    ws_summary.title = "Report Summary"
    for line in report_text.split("\n"):
        ws_summary.append([line])
    ws_summary.column_dimensions["A"].width = 70

    # 2. Events & Details
    ws_events = wb.create_sheet("Events & Details")
    ws_events.append([
        "Event #", "Location / Heading", "Description", "Status",
        "Sender", "Timestamp", "Source Message IDs", "Quoted Message ID", "Photo Available", "Photo"
    ])
    for idx, ev in enumerate(events, 1):
        ws_events.append([
            idx,
            ev.get("location_heading") or ev.get("level") or "N/A",
            ev.get("issue") or ev.get("activity") or ev.get("text") or "N/A",
            ev.get("status", "Pending"),
            ev.get("sender", "N/A"),
            format_report_timestamp(ev.get("datetime") or ev.get("raw_datetime")),
            ", ".join(ev.get("source_message_ids", [])),
            ev.get("quoted_message_id", "") or "None",
            "Yes" if (ev.get("photos") or ev.get("photos_data")) else "No",
            "",  # thumbnail drawn below
        ])
    apply_table_styling(ws_events)
    _embed_row_thumbnails(ws_events, events, col_letter="J", first_row=2)

    # 3. Needs Review / Unlinked
    ws_review = wb.create_sheet("Needs Review & Unlinked")
    ws_review.append(["Item ID", "Type / Reason", "Message Text", "Sender", "Timestamp"])
    for r in unlinked_or_needs_review:
        r = _review_row(r)
        ws_review.append([
            r.get("message_id", "N/A"),
            r.get("reason", "Unlinked reply"),
            r.get("text", "N/A"),
            r.get("sender", "N/A"),
            format_report_timestamp(r.get("datetime") or r.get("raw_datetime")),
        ])
    apply_table_styling(ws_review)

    # 4. Source Audit
    ws_audit = wb.create_sheet("Source Audit")
    ws_audit.append(["Message ID", "Group", "Sender", "Timestamp", "Quoted Message ID", "Audit Status"])
    for a in source_audit:
        ws_audit.append([
            a.get("message_id", "N/A"),
            a.get("group_name", "N/A"),
            a.get("sender", "N/A"),
            format_report_timestamp(a.get("timestamp") or a.get("raw_datetime")),
            a.get("quoted_message_id", "") or "Empty",
            a.get("audit_status", "Valid"),
        ])
    apply_table_styling(ws_audit)

    # 5. Photos Sheet (embed actual extracted images)
    all_photos = []
    for idx, ev in enumerate(events, 1):
        for p in ev.get("photos_data", []):
            all_photos.append((idx, ev, p))

    if all_photos:
        import io
        from openpyxl.drawing.image import Image as OpenpyxlImage
        ws_photos = wb.create_sheet("Photos")
        ws_photos.append(["Photo #", "Event #", "Location / Slot", "Description", "Image"])
        for p_idx, (ev_idx, ev, p_info) in enumerate(all_photos, 1):
            r_num = p_idx + 1
            ws_photos.append([
                p_idx,
                ev_idx,
                ev.get("location_heading") or ev.get("level") or "N/A",
                (ev.get("issue") or ev.get("activity") or ev.get("text") or "N/A")[:50],
                "",
            ])
            ws_photos.row_dimensions[r_num].height = 110
            img_data = p_info.get("bytes")
            if img_data:
                try:
                    img = OpenpyxlImage(io.BytesIO(img_data))
                    w, h = img.width, img.height
                    max_dim = 130
                    if w > max_dim or h > max_dim:
                        scale = min(max_dim / w, max_dim / h)
                        img.width = int(w * scale)
                        img.height = int(h * scale)
                    ws_photos.add_image(img, f"E{r_num}")
                except Exception:
                    pass
        ws_photos.column_dimensions["E"].width = 25
        apply_table_styling(ws_photos)

    out_file = OUTPUT_DIR / f"{filename_prefix}_{int(time.time())}.xlsx"
    wb.save(str(out_file))
    return out_file
