"""SQLite store of captured WhatsApp messages (written by the bot, see src/whatsapp/store.ts).

The reports engine reads it instead of re-parsing the whole Excel workbook on every report.
It is only used once the one-time Excel import has completed (meta key "excel_backfill_done"),
so a half-filled database can never make reports lose history.

    python3 -m reports_engine.messages_db            # import workspace/whatsapp_messages.xlsx
    python3 -m reports_engine.messages_db --check    # compare SQLite vs Excel without changing anything
"""

import argparse
import json
import os
import sqlite3
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

from .config_loader import BASE_DIR

EXCEL_FILE = BASE_DIR / "workspace" / "whatsapp_messages.xlsx"
BACKFILL_KEY = "excel_backfill_done"

# Keep in sync with MESSAGES_SCHEMA in src/whatsapp/store.ts
TABLE_DDL = """
    CREATE TABLE IF NOT EXISTS {table} (
        message_id        TEXT PRIMARY KEY,
        timestamp_text    TEXT NOT NULL,
        group_name        TEXT NOT NULL,
        sender            TEXT,
        text              TEXT,
        media_type        TEXT,
        quoted_message_id TEXT,
        thumb_jpeg        BLOB,
        media_file        TEXT,
        inserted_at       TEXT DEFAULT CURRENT_TIMESTAMP,
        ts_utc_ms         INTEGER
    )
"""
SCHEMA = TABLE_DDL.format(table="messages") + """;
    CREATE INDEX IF NOT EXISTS ix_messages_group ON messages(group_name);
    CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT);
"""

COLUMNS = "message_id, timestamp_text, group_name, sender, text, media_type, quoted_message_id, thumb_jpeg, media_file, ts_utc_ms"


def ensure_schema(c: sqlite3.Connection) -> None:
    """Create tables and add columns introduced after a database was first created."""
    c.executescript(SCHEMA)
    cols = {r[1] for r in c.execute("PRAGMA table_info(messages)")}
    if "ts_utc_ms" not in cols:
        c.execute("ALTER TABLE messages ADD COLUMN ts_utc_ms INTEGER")


def _ts_ms_from_text(text: Optional[str]) -> Optional[int]:
    from .date_utils import parse_datetime

    dt = parse_datetime(str(text or ""))
    return int(dt.timestamp() * 1000) if dt else None


def db_path_for(excel_path: Path) -> Path:
    """The bot keeps the database next to the workbook: whatsapp_messages.xlsx -> .db"""
    return excel_path.with_suffix(".db")


def is_ready(db_path: Path) -> bool:
    if not db_path.exists():
        return False
    try:
        with sqlite3.connect(str(db_path)) as c:
            return c.execute("SELECT 1 FROM meta WHERE key=?", (BACKFILL_KEY,)).fetchone() is not None
    except sqlite3.Error:
        return False


def source_mtime(db_path: Path) -> float:
    """Change marker for caching: WAL mode writes to the -wal file first."""
    marks = [p.stat().st_mtime for p in (db_path, Path(str(db_path) + "-wal")) if p.exists()]
    return max(marks) if marks else 0.0


def read_rows(
    db_path: Path, load_images: bool
) -> Tuple[List[Tuple[int, tuple]], Dict[int, List[Dict[str, Any]]], Dict[int, int]]:
    """
    Rows shaped like the Excel sheet (id, time, group, sender, text, type, photo preview, quoted id)
    so the parser treats both sources identically. Row numbers are synthetic (2..n+1).
    Also returns the exact message time (UTC epoch ms) per row where the bot recorded it.
    """
    rows: List[Tuple[int, tuple]] = []
    images: Dict[int, List[Dict[str, Any]]] = {}
    ts_by_row: Dict[int, int] = {}
    thumb_col = "thumb_jpeg" if load_images else "NULL"
    with sqlite3.connect(str(db_path)) as c:
        cols = {r[1] for r in c.execute("PRAGMA table_info(messages)")}
        ts_col = "ts_utc_ms" if "ts_utc_ms" in cols else "NULL"
        cur = c.execute(
            f"SELECT message_id, timestamp_text, group_name, sender, text, media_type, quoted_message_id, {thumb_col}, {ts_col} "
            "FROM messages ORDER BY rowid"
        )
        for idx, (mid, ts, grp, sender, text, mtype, quoted, thumb, ts_ms) in enumerate(cur, start=2):
            rows.append((idx, (mid, ts, grp, sender, text, mtype, "", quoted)))
            if thumb:
                images[idx] = [{"media_path": f"sqlite:{mid}", "bytes": bytes(thumb)}]
            if ts_ms:
                ts_by_row[idx] = int(ts_ms)
    return rows, images, ts_by_row


def _excel_rows(excel_path: Path) -> List[Tuple[int, tuple, Optional[bytes]]]:
    from .whatsapp_parser import extract_xlsx_images

    images = extract_xlsx_images(excel_path)
    wb = openpyxl.load_workbook(str(excel_path), data_only=True, read_only=True)
    out = []
    for row_idx, row in enumerate(wb.active.iter_rows(min_row=2, values_only=True), start=2):
        if row and len(row) >= 5 and row[0] and row[2]:
            pics = images.get(row_idx) or []
            out.append((row_idx, row, pics[0]["bytes"] if pics else None))
    wb.close()
    return out


def import_excel(excel_path: Path = EXCEL_FILE, db_path: Optional[Path] = None) -> Dict[str, int]:
    """
    Rebuild the messages table in Excel order, keep any newer rows the bot already stored,
    then mark the database ready. Safe to re-run.
    """
    db_path = db_path or db_path_for(excel_path)
    rows = _excel_rows(excel_path)
    with sqlite3.connect(str(db_path)) as c:
        c.execute("PRAGMA journal_mode=WAL")
        ensure_schema(c)
        c.execute("DROP TABLE IF EXISTS messages_new")
        c.execute(TABLE_DDL.format(table="messages_new"))
        for _, row, thumb in rows:
            cell = lambda i: row[i] if len(row) > i else None
            c.execute(
                f"INSERT OR IGNORE INTO messages_new ({COLUMNS}) VALUES (?,?,?,?,?,?,?,?,?,?)",
                (
                    str(cell(0)).strip(),
                    str(cell(1) or ""),
                    str(cell(2)).strip(),
                    str(cell(3) or ""),
                    str(cell(4) or ""),
                    str(cell(5) or "") if len(row) > 5 else "Text",
                    (str(cell(7)) if cell(7) is not None else None) if len(row) > 7 else None,
                    thumb,
                    None,
                    None,  # filled below: bot's exact time if it has the row, else parsed from the text
                ),
            )
        # Rows the bot stored that the workbook does not have yet; keep their thumbnails too
        bot_only = c.execute(
            f"INSERT OR IGNORE INTO messages_new ({COLUMNS}) SELECT {COLUMNS} FROM messages ORDER BY rowid"
        ).rowcount
        c.execute(
            "UPDATE messages_new SET thumb_jpeg = (SELECT m.thumb_jpeg FROM messages m WHERE m.message_id = messages_new.message_id) "
            "WHERE thumb_jpeg IS NULL"
        )
        c.execute(
            "UPDATE messages_new SET ts_utc_ms = (SELECT m.ts_utc_ms FROM messages m WHERE m.message_id = messages_new.message_id) "
            "WHERE ts_utc_ms IS NULL"
        )
        missing_ts = c.execute("SELECT message_id, timestamp_text FROM messages_new WHERE ts_utc_ms IS NULL").fetchall()
        c.executemany(
            "UPDATE messages_new SET ts_utc_ms = ? WHERE message_id = ?",
            [(_ts_ms_from_text(text), mid) for mid, text in missing_ts],
        )
        c.execute("DROP TABLE messages")
        c.execute("ALTER TABLE messages_new RENAME TO messages")
        c.execute("CREATE INDEX IF NOT EXISTS ix_messages_group ON messages(group_name)")
        total = c.execute("SELECT COUNT(*) FROM messages").fetchone()[0]
        with_photo = c.execute("SELECT COUNT(*) FROM messages WHERE thumb_jpeg IS NOT NULL").fetchone()[0]
        c.execute(
            "INSERT OR REPLACE INTO meta (key, value) VALUES (?, ?)",
            (BACKFILL_KEY, datetime.now(timezone.utc).isoformat()),
        )
        c.commit()
    return {"excel_rows": len(rows), "bot_only_rows": bot_only, "total": total, "with_photo": with_photo}


def fill_missing_ts(db_path: Path) -> int:
    """Give rows stored before ts_utc_ms existed a time parsed from their text. Returns rows updated."""
    with sqlite3.connect(str(db_path)) as c:
        ensure_schema(c)
        missing = c.execute("SELECT message_id, timestamp_text FROM messages WHERE ts_utc_ms IS NULL").fetchall()
        c.executemany("UPDATE messages SET ts_utc_ms = ? WHERE message_id = ?",
                      [(_ts_ms_from_text(text), mid) for mid, text in missing])
        c.commit()
    return len(missing)


# Picture box in the messages workbook: the 2.30" Excel thumbnail width, taller for upright videos
PICTURE_MAX_W, PICTURE_MAX_H = 221, 260
# The download button sends this many recent days; --days 0 exports everything
EXPORT_DEFAULT_DAYS = 30
HEADERS = ["Message ID", "Date & Time", "Group Name", "Sender", "Message", "Type", "Photo Preview",
           "Quoted Message ID", "Time (Dubai)", "Reply to", "Ticket", "Ticket status"]


def _ticket_status(tickets_path: Optional[Path]) -> Dict[str, Tuple[str, str]]:
    """message id -> (report, status) for the messages that became maintenance / fire safety tickets."""
    if not tickets_path or not Path(tickets_path).exists():
        return {}
    try:
        with sqlite3.connect(str(tickets_path)) as c:
            return {r[0]: (r[1] or "", r[2] or "") for r in c.execute("SELECT message_id, report_type, status FROM tickets")}
    except sqlite3.Error:
        return {}


def export_excel(out_path: Path, db_path: Path, days: Optional[int] = None,
                 tickets_path: Optional[Path] = None) -> int:
    """
    Build the WhatsApp messages workbook from the SQLite store, in message-time order, for reading:
    names without numbers, @mentions as names, a sortable Dubai time, the text of the message replied
    to, and the ticket each message became. Columns A-H keep the live workbook's layout.
    days limits it to the most recent days (None = everything). Used by the Telegram download button.
    """
    import io as _io
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.styles import Alignment, Font, PatternFill

    from .date_utils import DUBAI_TZ
    from .master_data import display_name
    from .sender_resolver import MENTION_REGEX, SenderResolver

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WhatsApp Messages"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E293B")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for col, width in zip("ABCDEFGHIJKL", (36, 22, 28, 26, 70, 12, 28, 36, 18, 45, 14, 14)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    since_ms = int((datetime.now(timezone.utc).timestamp() - days * 86400) * 1000) if days else None
    with sqlite3.connect(str(db_path)) as c:
        ensure_schema(c)
        resolver = SenderResolver()
        for (sender,) in c.execute("SELECT DISTINCT sender FROM messages"):
            resolver.register_sender(sender or "")
        where = "WHERE ts_utc_ms >= ?" if since_ms else ""
        rows = c.execute(
            "SELECT message_id, timestamp_text, group_name, sender, text, media_type, quoted_message_id, thumb_jpeg, ts_utc_ms "
            f"FROM messages {where} ORDER BY COALESCE(ts_utc_ms, 0), rowid", (since_ms,) if since_ms else ()
        ).fetchall()
        quoted_ids = {r[6] for r in rows if r[6]}
        quoted_text: Dict[str, str] = {}
        for i in range(0, len(quoted_ids), 900):
            chunk = list(quoted_ids)[i:i + 900]
            for mid, sender, text in c.execute(
                    f"SELECT message_id, sender, text FROM messages WHERE message_id IN ({','.join('?' * len(chunk))})", chunk):
                quoted_text[mid] = f"{display_name(sender)}: {' '.join((text or '').split())[:120]}"
    tickets = _ticket_status(tickets_path)

    def readable(text: str) -> str:
        # Mentions the store cannot name stay anonymous: the id can be a phone number
        return MENTION_REGEX.sub("@someone", resolver.resolve_mentions(text or ""))

    wrap = Alignment(wrap_text=True, vertical="center")
    for i, (mid, ts, grp, sender, text, mtype, quoted, thumb, ts_ms) in enumerate(rows, start=2):
        local = datetime.fromtimestamp(ts_ms / 1000, tz=DUBAI_TZ).replace(tzinfo=None) if ts_ms else None
        is_media = any(k in (mtype or "") for k in ("Image", "Video"))
        note = "" if thumb or not is_media else (
            "No picture in the chat export" if str(mid).startswith("EXP") else "No picture saved")
        if quoted:
            reply_to = quoted_text.get(quoted, "(older message, not recorded)")
        else:
            reply_to = ""
        report, status = tickets.get(mid, ("", ""))
        ws.append([mid, ts, grp, display_name(sender), readable(text), mtype, note, quoted or "",
                   local, readable(reply_to), report.replace("_", " ").title(), status])
        ws.cell(row=i, column=5).alignment = wrap
        ws.cell(row=i, column=10).alignment = wrap
        ws.cell(row=i, column=9).number_format = "yyyy-mm-dd hh:mm"
        if thumb:
            try:
                img = XlImage(_io.BytesIO(bytes(thumb)))
                # Video frames are stored larger than the cell: scale down to fit it
                scale = min(1.0, PICTURE_MAX_W / img.width, PICTURE_MAX_H / img.height)
                img.width, img.height = int(img.width * scale), int(img.height * scale)
                ws.add_image(img, f"G{i}")
                ws.row_dimensions[i].height = max(img.height * 0.75 + 4, 20)
            except Exception:
                pass
    ws.auto_filter.ref = f"A1:L{max(len(rows) + 1, 1)}"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return len(rows)


def purge_hr(db_path: Path, retention_days: Optional[int], markers: List[str]) -> int:
    """Delete HR / leave messages older than retention_days from the local store (None = keep)."""
    if retention_days is None or not db_path.exists() or not markers:
        return 0
    cutoff_ms = int((datetime.now(timezone.utc).timestamp() - retention_days * 86400) * 1000)
    where = " OR ".join("lower(text) LIKE ?" for _ in markers)
    with sqlite3.connect(str(db_path)) as c:
        ensure_schema(c)
        n = c.execute(
            f"DELETE FROM messages WHERE ts_utc_ms IS NOT NULL AND ts_utc_ms < ? AND ({where})",
            [cutoff_ms] + [f"%{m.lower()}%" for m in markers],
        ).rowcount
        c.commit()
    return n


def compare(excel_path: Path = EXCEL_FILE, db_path: Optional[Path] = None) -> Dict[str, Any]:
    """Parse both sources with the real parser and report differences (no writes)."""
    from .whatsapp_parser import WhatsAppParser

    db_path = db_path or db_path_for(excel_path)
    prev = os.environ.get("REPORTS_SOURCE")
    try:
        os.environ["REPORTS_SOURCE"] = "excel"
        from_excel = {m["message_id"]: m for m in WhatsAppParser(excel_path).load_messages(load_images=True)}
        os.environ["REPORTS_SOURCE"] = "sqlite"
        from_db = {m["message_id"]: m for m in WhatsAppParser(excel_path).load_messages(load_images=True)}
    finally:
        if prev is None:
            os.environ.pop("REPORTS_SOURCE", None)
        else:
            os.environ["REPORTS_SOURCE"] = prev
    fields = ("raw_datetime", "datetime", "group_name", "sender", "text", "media_type", "quoted_message_id", "quoted_captured")
    diffs = [
        mid for mid in from_excel.keys() & from_db.keys()
        if any(from_excel[mid][f] != from_db[mid][f] for f in fields)
        or len(from_excel[mid]["photos_data"]) != len(from_db[mid]["photos_data"])
    ]
    return {
        "excel": len(from_excel),
        "sqlite": len(from_db),
        "missing_in_sqlite": len(from_excel.keys() - from_db.keys()),
        "only_in_sqlite": len(from_db.keys() - from_excel.keys()),
        "field_differences": len(diffs),
        "ready": is_ready(db_path),
    }


def main() -> int:
    ap = argparse.ArgumentParser(description="Import the WhatsApp Excel into the SQLite message store")
    ap.add_argument("--excel", default=str(EXCEL_FILE))
    ap.add_argument("--db", default=None)
    ap.add_argument("--check", action="store_true", help="Only compare SQLite with Excel")
    ap.add_argument("--fill-ts", action="store_true", help="Only fill missing exact timestamps from the text")
    ap.add_argument("--export-excel", metavar="OUT", help="Write the messages workbook built from SQLite to OUT")
    ap.add_argument("--days", type=int, default=EXPORT_DEFAULT_DAYS,
                    help="with --export-excel: only the last N days (0 = every message)")
    args = ap.parse_args()
    excel, db = Path(args.excel), Path(args.db) if args.db else None
    if args.export_excel:
        target = db or db_path_for(excel)
        if not is_ready(target):
            print("SQLite store not ready", file=sys.stderr)
            return 2
        from .tickets_db import DEFAULT_DB_PATH as TICKETS_DB
        n = export_excel(Path(args.export_excel), target, days=args.days or None, tickets_path=Path(TICKETS_DB))
        print(json.dumps({"exported": n, "path": str(Path(args.export_excel).resolve())}))
        return 0
    if args.fill_ts:
        target = db or db_path_for(excel)
        n = fill_missing_ts(target) if target.exists() else 0
        print(f"✅ Filled exact time for {n} older message(s)")
        return 0
    if not excel.exists():
        print(f"Excel not found: {excel}", file=sys.stderr)
        return 1
    if not args.check:
        r = import_excel(excel, db)
        print(f"✅ Imported {r['excel_rows']} Excel rows (+{r['bot_only_rows']} newer rows from the bot) "
              f"-> {r['total']} messages, {r['with_photo']} with photo")
    c = compare(excel, db)
    ok = c["missing_in_sqlite"] == 0 and c["field_differences"] == 0 and c["ready"]
    print(f"{'✅' if ok else '❌'} Excel {c['excel']} | SQLite {c['sqlite']} | missing {c['missing_in_sqlite']} | "
          f"newer in SQLite {c['only_in_sqlite']} | field differences {c['field_differences']} | ready {c['ready']}")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
